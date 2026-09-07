"""Resolve every cross-sheet reference to the LABEL it lands on.

LibreOffice cannot recalculate in this environment, so this is the check that
matters more anyway: a formula can evaluate perfectly and point at the wrong row.
"""
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
import re, collections
wb = load_workbook("/home/user/etablix/tools/village-plan.xlsx")
REF = re.compile(r"(?:'([^']+)'|([A-Za-z][A-Za-z0-9_]*))!\$?([A-Z]{1,2})\$?(\d+)")
BANNED = re.compile(r"\b(XLOOKUP|XMATCH|SORT|FILTER|UNIQUE|SEQUENCE|TEXTJOIN|CONCAT|IFS|SWITCH|MAXIFS|MINIFS)\s*\(", re.I)

def label_at(sheet, row):
    ws = wb[sheet]
    for c in (2, 3, 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()[:52]
    return "(no label)"

seen, banned, bad = collections.Counter(), [], []
uniq = {}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.startswith("="):
                continue
            if BANNED.search(v):
                banned.append((ws.title, cell.coordinate, v[:70]))
            for m in REF.finditer(v):
                tgt = m.group(1) or m.group(2)
                if tgt not in wb.sheetnames:
                    bad.append((ws.title, cell.coordinate, tgt)); continue
                r = int(m.group(4))
                key = (ws.title, tgt, m.group(3), r)
                seen[key] += 1
                uniq.setdefault((ws.title, tgt, m.group(3), r), cell.coordinate)

print("=== CROSS-SHEET REFERENCES — what each one actually lands on ===\n")
by_src = collections.defaultdict(list)
for (src, tgt, c, r), n in seen.items():
    if src == tgt: continue
    by_src[src].append((tgt, c, r, n))
for src in sorted(by_src):
    print(f"── {src}")
    for tgt, c, r, n in sorted(set(by_src[src])):
        print(f"     → {tgt}!{c}{r:<4} ×{n:<4} {label_at(tgt, r)}")
    print()
print("unknown sheet references:", len(bad), bad[:5])
print("banned/unsupported functions:", len(banned), banned[:5])
