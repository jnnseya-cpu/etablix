"""Weightings within quality sum to 100%; the quality/price split is separate."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
exec(open("styles.py").read())
E = json.load(open("eval.json")); E2 = json.load(open("eval2.json"))
Q0, Q1, QT, PW = E["Q0"], E["Q1"], E["QT"], E["PW"]
wb = load_workbook("/home/user/etablix/tools/tender-evaluation.xlsx")
ws = wb["Criteria"]

# criteria weightings, as a share of the quality element — these total 100%
NEW = [0.19, 0.19, 0.15, 0.10, 0.12, 0.10, 0.09, 0.06]
assert abs(sum(NEW) - 1) < 1e-9, sum(NEW)
for i, w in enumerate(NEW):
    c = ws.cell(row=Q0 + i, column=4, value=w); c.font = INPUT; c.number_format = "0.0%"; c.fill = KEYFILL
ws.cell(row=5, column=4, value="Weighting\nwithin quality").alignment = Alignment(wrap_text=True, horizontal="center", vertical="bottom")
ws.cell(row=QT, column=3, value="QUALITY CRITERIA — TOTAL (must be 100%)").font = H3
ws.cell(row=QT, column=5, value="These are shares of the quality element, not of the whole evaluation. The quality/price split is set below.").font = SMALL

# replace the old price row with an explicit split
ws.cell(row=PW, column=2, value="")
ws.cell(row=PW, column=3, value="QUALITY / PRICE SPLIT — quality share").font = H3
c = ws.cell(row=PW, column=4, value=0.60); c.font = INPUT; c.number_format = "0.0%"; c.fill = KEYFILL
ws.cell(row=PW, column=5, value="The split is a decision, not a default. Where the risk is interface failure rather than rate, a higher quality "
        "share buys a better outcome — and you must be able to say why you chose it, before tenders are opened.").font = SMALL
ws.cell(row=PW, column=5).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[PW].height = 30
PS = PW + 1
ws.cell(row=PS, column=3, value="                                                    price share").font = H3
ws.cell(row=PS, column=4, value=f"=1-D{PW}").font = FORMULA
ws.cell(row=PS, column=4).number_format = "0.0%"; ws.cell(row=PS, column=4).fill = TOTFILL
ws.cell(row=PS, column=5, value="Follows from the quality share. The two always total 100%.").font = SMALL
ws.cell(row=PS + 1, column=3, value="Check").font = H3
ws.cell(row=PS + 1, column=4,
    value=f'=IF(ROUND(D{QT},4)<>1,"ERROR: quality criteria total "&TEXT(D{QT},"0.0%")&" — they must total 100%",'
          f'IF(ROUND(D{PW}+D{PS},4)<>1,"ERROR: the split does not total 100%","Correct — criteria total 100%, split totals 100%"))')
ws.cell(row=PS + 1, column=4).font = Font(name="Arial", size=10, bold=True, color="C0392B")
ws.merge_cells(start_row=PS + 1, start_column=4, end_row=PS + 1, end_column=5)
ws.cell(row=PS + 2, column=4, value=None)

# Scores: weighted quality is now a straight SUMPRODUCT (criteria already total 100%)
sc = wb["Scores"]; QSC = E2["QSCORE"]; S0, S1 = E2["S0"], E2["S1"]
TC = [get_column_letter(4 + i) for i in range(5)]
for j, c in enumerate(TC):
    sc.cell(row=QSC, column=4 + j, value=f"=SUMPRODUCT({c}{S0}:{c}{S1},Criteria!$D${Q0}:$D${Q1})")
# Result: combine on the split
rs = wb["Result"]
COMB = 10
for i, c in enumerate(TC):
    rs.cell(row=COMB, column=4 + i, value=f"={c}7*Criteria!$D${PW}+{c}8*Criteria!$D${PS}")
# Sensitivity band already uses (1-pw) and pw directly — correct and independent of the split
json.dump({"PS": PS}, open("eval3.json", "w"))
wb.save("/home/user/etablix/tools/tender-evaluation.xlsx")
print(f"criteria now total {sum(NEW):.0%}; quality/price split 60/40 at rows {PW}/{PS}")
