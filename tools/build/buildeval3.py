"""Result + Sensitivity — the sheet that decides whether the award survives."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
exec(open("styles.py").read())
E = json.load(open("eval.json")); E2 = json.load(open("eval2.json"))
QT, PW = E["QT"], E["PW"]; PSC, QSC, TOTROW = E2["PSCORE"], E2["QSCORE"], E2["TOTROW"]
GBP = '£#,##0;(£#,##0);-'
wb = load_workbook("/home/user/etablix/tools/tender-evaluation.xlsx")
N = 5; TC = [get_column_letter(4 + i) for i in range(N)]

ws = wb.create_sheet("Result"); ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 6; ws.column_dimensions["C"].width = 42
for c in TC: ws.column_dimensions[c].width = 16
ws.column_dimensions[get_column_letter(4 + N + 1)].width = 54
ws["B2"] = "RESULT"; ws["B2"].font = H1
ws["B3"] = "Quality and price combined on the published split. Then the question that decides whether this award survives being challenged."; ws["B3"].font = SMALL
ws.cell(row=5, column=3, value="Tenderer").font = H3
for i, c in enumerate(TC):
    cell = ws.cell(row=5, column=4 + i, value=f"=Prices!{c}5"); cell.font = LINK
    cell.alignment = Alignment(horizontal="center"); cell.fill = SECFILL

rows = [
 ("Tender total", f"=Prices!{{c}}{TOTROW}", GBP, ""),
 ("Quality score (of 5)", f"=Scores!{{c}}{QSC}", "0.00", ""),
 ("Price score (of 5)", f"=Prices!{{c}}{PSC}", "0.00", ""),
]
r = 6
for lab, f, fmt, note in rows:
    ws.cell(row=r, column=3, value=lab).font = BODY
    for i, c in enumerate(TC):
        cell = ws.cell(row=r, column=4 + i, value=f.replace("{c}", c))
        cell.font = FORMULA; cell.number_format = fmt; cell.alignment = Alignment(horizontal="center")
    r += 1
r += 1
ws.cell(row=r, column=3, value="COMBINED SCORE (of 5)").font = H3
for i, c in enumerate(TC):
    cell = ws.cell(row=r, column=4 + i, value=f"={c}7*Criteria!$D${QT}+{c}8*Criteria!$D${PW}")
    cell.font = H3; cell.number_format = "0.000"; cell.fill = TOTFILL
    cell.alignment = Alignment(horizontal="center")
COMB = r; r += 1
ws.cell(row=r, column=3, value="Rank").font = H3
for i, c in enumerate(TC):
    cell = ws.cell(row=r, column=4 + i, value=f"=IF({c}6=0,\"\",RANK({c}{COMB},$D${COMB}:${TC[-1]}${COMB},0))")
    cell.font = H3; cell.alignment = Alignment(horizontal="center")
RANK = r; r += 2

facts = [
 ("Winning tenderer", f'=IFERROR(INDEX($D$5:${TC[-1]}$5,MATCH(1,$D${RANK}:${TC[-1]}${RANK},0)),"—")', ""),
 ("Winning combined score", f"=MAX($D${COMB}:${TC[-1]}${COMB})", "The number the award is made on."),
 ("Second place score", f"=LARGE($D${COMB}:${TC[-1]}${COMB},2)", ""),
 ("Margin, first over second", f"=D{r+1}-D{r+2}", "The whole question is whether this margin is bigger than the noise in the scoring."),
 ("Margin as % of the winning score", f"=IFERROR(D{r+3}/D{r+1},0)", "Under about 3% the two tenders are, honestly, level. Say so and decide on something you can defend rather than on a decimal."),
]
for lab, f, note in facts:
    ws.cell(row=r, column=3, value=lab).font = BODY
    cell = ws.cell(row=r, column=4, value=f); cell.font = H3
    cell.number_format = "0.0%" if "%" in lab else ("0.000" if "score" in lab.lower() or "Margin" in lab else "General")
    if "Winning tenderer" in lab: cell.number_format = "General"
    n = ws.cell(row=r, column=4 + N + 1, value=note); n.font = SMALL
    n.alignment = Alignment(wrap_text=True, vertical="top")
    if len(note) > 60: ws.row_dimensions[r].height = 28
    r += 1
MARGINPC = r - 1
r += 1
band(ws, r, "   SENSITIVITY — WOULD A DIFFERENT WEIGHTING HAVE CHANGED THE WINNER?", 8); r += 1
ws.cell(row=r, column=3, value=("The weightings were published before tenders were opened and are not changed. This is not a re-run — it is a "
        "test of how much the result depends on a judgement made before anyone knew the prices. If the winner flips inside a reasonable band, "
        "the award is fragile, and it is far better to know that now than in a letter from the runner-up's solicitor.")).font = Font(name="Arial", size=9, italic=True)
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4 + N + 1)
ws.row_dimensions[r].height = 32; r += 1
ws.cell(row=r, column=3, value="Price weighting tested").font = H3
for i, c in enumerate(TC):
    ws.cell(row=r, column=4 + i, value=f"=Prices!{c}5").font = LINK
    ws.cell(row=r, column=4 + i).alignment = Alignment(horizontal="center"); ws.cell(row=r, column=4 + i).fill = SECFILL
ws.cell(row=r, column=4 + N + 1, value="Winner at this weighting").font = H3
ws.cell(row=r, column=4 + N + 1).fill = SECFILL
SH = r; r += 1
for pct in (0.20, 0.30, 0.40, 0.50, 0.60, 0.70):
    ws.cell(row=r, column=3, value=pct).font = INPUT
    ws.cell(row=r, column=3).number_format = "0%"
    for i, c in enumerate(TC):
        cell = ws.cell(row=r, column=4 + i, value=f"={c}7*(1-$C{r})+{c}8*$C{r}")
        cell.font = FORMULA; cell.number_format = "0.000"; cell.alignment = Alignment(horizontal="center")
    w = ws.cell(row=r, column=4 + N + 1,
        value=f'=IFERROR(INDEX($D${SH}:${TC[-1]}${SH},MATCH(MAX($D{r}:${TC[-1]}{r}),$D{r}:${TC[-1]}{r},0)),"—")')
    w.font = H3
    r += 1
SB0, SB1 = SH + 1, r - 1
r += 1
ws.cell(row=r, column=3, value="Winner stable across the band?").font = H3
ws.cell(row=r, column=4, value=f'=IF(COUNTIF({get_column_letter(4+N+1)}{SB0}:{get_column_letter(4+N+1)}{SB1},{get_column_letter(4+N+1)}{SB0})=6,'
        f'"STABLE — the same tenderer wins from 20% to 70% price weighting","FRAGILE — the winner changes with the weighting")')
ws.cell(row=r, column=4).font = Font(name="Arial", size=11, bold=True, color="C0392B")
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=4 + N + 1)
r += 2
ws.cell(row=r, column=3, value=("A FRAGILE result is not a reason to change the weighting. It is a reason to look harder at the quality scores, "
        "moderate them properly, and be certain the reasons written against each one would stand up if read aloud. If the result is still "
        "fragile after that, the two tenders really are equivalent — and the honest answer is to say so and decide on a stated tie-break, "
        "not to pretend a decimal place settled it.")).font = Font(name="Arial", size=9, italic=True, color="C0392B")
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4 + N + 1)
ws.row_dimensions[r].height = 44
wb.save("/home/user/etablix/tools/tender-evaluation.xlsx")
print("Result + Sensitivity written. Combined row", COMB, "sensitivity band rows", SB0, "-", SB1)
