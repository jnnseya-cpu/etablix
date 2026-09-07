"""Recompute the Packaging sheet independently."""
from openpyxl import load_workbook
import json
CAP = json.load(open("cap2.json")); pk0 = CAP["PK0"]
wb = load_workbook("/home/user/etablix/tools/village-plan.xlsx")
ws = wb["Packaging"]
import re
def g(cell):
    v = ws[cell].value
    # a prime-column cell that simply mirrors the five-package column
    while isinstance(v, str) and re.fullmatch(r"=[A-D]\d+", v):
        v = ws[v[1:]].value
    return v
# inputs
c5, c1 = g("C7"), g("D7")
if5, if1 = c5*(c5-1)/2, c1*(c1-1)/2
# direct cost from the recompute
exec(open("recompute.py").read().split("P = lambda")[0])
direct = sum(cap.values())          # P1..P5 capex
directop = opex
mg5 = g("C12")*g("C13")*g("C14"); mg1 = g("D12")*g("D13")*g("D14")
ifc5 = if5*g("C17")*g("C18"); ifc1 = if5*g("C17")*g("D18")   # same interface count, lower leak
mar5 = (direct+directop)*g("C21"); mar1 = (direct+directop)*g("D21")
pr5, pr1 = g("C23"), g("D23")
t5 = direct+directop+mg5+ifc5+mar5+pr5
t1 = direct+directop+mg1+ifc1+mar1+pr1
be = (t5-direct-directop-mg1-ifc1-pr1)/(direct+directop)
P = lambda l,a,b="": print(f"  {l:<38}{a:>16}{b:>16}")
print("                                        FIVE PACKAGES        ONE PRIME")
P("Contracts", c5, c1)
P("Interfaces the client owns", int(if5), int(if1))
P("Direct capex + opex", f"£{direct+directop:,.0f}", f"£{direct+directop:,.0f}")
P("Client management", f"£{mg5:,.0f}", f"£{mg1:,.0f}")
P("Expected interface cost", f"£{ifc5:,.0f}", f"£{ifc1:,.0f}")
P("Prime margin", f"£{mar5:,.0f}", f"£{mar1:,.0f}")
P("Procurement cost", f"£{pr5:,.0f}", f"£{pr1:,.0f}")
P("TOTAL", f"£{t5:,.0f}", f"£{t1:,.0f}")
print()
P("Difference (positive = prime cheaper)", f"£{t5-t1:,.0f}")
P("Break-even prime margin", f"{be:.2%}")
print(f"\n  At the {g('D21'):.1%} margin assumed, one prime is {'cheaper' if t1<t5 else 'DEARER'} by £{abs(t5-t1):,.0f}.")
print(f"  The prime stops paying above a margin of {be:.2%}.")
