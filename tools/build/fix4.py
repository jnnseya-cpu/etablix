from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
exec(open("styles.py").read())
wb = load_workbook("/home/user/etablix/tools/village-plan.xlsx")
ws = wb["BedDemand"]
# helper row 15: month number where there is demand, blank otherwise
ws.cell(row=15, column=2, value="Helper — month number where rooms are needed").font = SMALL
for m in range(1, MONTHS + 1):
    cl = col(C0 + m - 1)
    c = ws.cell(row=15, column=C0 + m - 1, value=f'=IF({cl}14>0,{cl}$6,"")')
    c.font = SMALL; c.alignment = Alignment(horizontal="center")
ws["C23"] = f"=MIN({col(C0)}15:{col(CN)}15)"
ws["C24"] = f"=MAX({col(C0)}15:{col(CN)}15)"
for cell in ("C23", "C24"):
    ws[cell].font = FORMULA; ws[cell].number_format = "0"
wb.save("/home/user/etablix/tools/village-plan.xlsx")
print("array formulas replaced with a helper row")
