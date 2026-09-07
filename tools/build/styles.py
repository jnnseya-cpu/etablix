from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
INPUT = Font(name="Arial", size=10, color="0000FF")
FORMULA = Font(name="Arial", size=10)
LINK = Font(name="Arial", size=10, color="008000")
H1 = Font(name="Arial", size=14, bold=True, color="14181D")
H2 = Font(name="Arial", size=11, bold=True, color="FFFFFF")
H3 = Font(name="Arial", size=10, bold=True)
BODY = Font(name="Arial", size=10)
SMALL = Font(name="Arial", size=9, color="5B6672")
KEYFILL = PatternFill("solid", fgColor="FFF9C4")
BANDFILL = PatternFill("solid", fgColor="14181D")
SECFILL = PatternFill("solid", fgColor="EDE7DA")
WARNFILL = PatternFill("solid", fgColor="FDEDEC")
TOTFILL = PatternFill("solid", fgColor="F2EFE7")
THIN = Side(style="thin", color="D5D5D5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONTHS = 38; C0 = 3; CN = C0 + MONTHS - 1
def col(i): return get_column_letter(i)
def band(ws, row, text, width=14):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row=row, column=1, value=text); c.font = H2; c.fill = BANDFILL
    c.alignment = Alignment(vertical="center"); ws.row_dimensions[row].height = 20
def section(ws, row, text, width=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row=row, column=1, value=text); c.font = H3; c.fill = SECFILL
