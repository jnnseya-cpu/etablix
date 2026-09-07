"""Packaging sheet — five contracts against one prime, modelled."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
exec(open("styles.py").read())
A = json.load(open("assump_rows.json")); CAP = json.load(open("cap2.json"))
wb = load_workbook("/home/user/etablix/tools/village-plan.xlsx")
if "Packaging" in wb.sheetnames: del wb["Packaging"]
GBP = '£#,##0;(£#,##0);-'
pk0 = CAP["PK0"]
OPX = f"Opex!$AP$14"

ws = wb.create_sheet("Packaging"); ws.sheet_view.showGridLines = False
for c, w in zip("ABCDEFG", (4, 48, 15, 15, 13, 56, 2)): ws.column_dimensions[c].width = w
ws["B2"] = "FIVE CONTRACTS, OR ONE — MODELLED"; ws["B2"].font = H1
ws["B3"] = ("Splitting the village into five packages buys competition on each. Letting it as one buys an owner for the whole. "
            "This sheet prices the difference instead of arguing about it."); ws["B3"].font = SMALL

K = {}
r = 5
def sec(t):
    global r
    section(ws, r, t, 6); r += 1
def line(key, label, five, one, unit, basis, fmt=GBP, bold=False, kind="formula"):
    global r
    # Register first, so a formula on this row can reference its own row.
    if key: K[key] = r
    c = ws.cell(row=r, column=2, value=label); c.font = H3 if bold else BODY
    for cc, val in ((3, five), (4, one)):
        if isinstance(val, str) and val.startswith("="):
            val = val.format(**{k: str(vv) for k, vv in K.items()})
        v = ws.cell(row=r, column=cc, value=val)
        v.font = H3 if bold else (INPUT if kind == "input" else FORMULA)
        v.number_format = fmt
        if bold: v.fill = TOTFILL
        if kind == "input": v.fill = KEYFILL
    ws.cell(row=r, column=5, value=unit).font = SMALL
    b = ws.cell(row=r, column=6, value=basis); b.font = SMALL
    b.alignment = Alignment(wrap_text=True, vertical="top")
    if len(basis) > 72: ws.row_dimensions[r].height = 26
    for cc in range(2, 7): ws.cell(row=r, column=cc).border = BOX
    r += 1

for i, h in enumerate(["", "", "FIVE PACKAGES", "ONE PRIME", "Unit", "Basis — and what to test"]):
    c = ws.cell(row=r, column=i + 1, value=h); c.font = H3
    c.alignment = Alignment(horizontal="center" if i in (2, 3) else "left")
    c.border = Border(bottom=Side(style="medium", color="14181D"))
r += 1

sec("WHAT YOU ARE BUYING")
line("contracts", "Contracts the client holds", 5, 1, "no.", "Change the left-hand figure if you split differently. Everything below follows it.", fmt="0", kind="input")
line("ifaces", "Interfaces the client owns", "=C{contracts}*(C{contracts}-1)/2", "=D{contracts}*(D{contracts}-1)/2", "no.",
     "n(n−1)/2. Five contracts is ten interfaces; one is none — they move inside the prime, they do not disappear. The Packages sheet lists fifteen because several involve three parties.", fmt="0")
line("direct", "Direct package cost — capex", f"=Capex!$D${pk0}+Capex!$D${pk0+1}+Capex!$D${pk0+2}+Capex!$D${pk0+3}+Capex!$D${pk0+4}",
     "=C{direct}", "£", "The same physical village either way. Consolidation does not change what gets built.")
line("directop", "Direct package cost — opex over the programme", f"={OPX}", "=C{directop}", "£", "Same.")

sec("WHAT THE CLIENT SPENDS TO MANAGE IT")
line("mgfte", "Client-side management FTE", 2.5, 0.8, "FTE", "Package managers, commercial and coordination on the client's own payroll. One prime is not zero — somebody still holds the prime to account.", fmt="0.0", kind="input")
line("mgrate", "Loaded cost per FTE per month", 9500, "=C{mgrate}", "£", "INDICATIVE. Replace with your own.", kind="input")
line("months", "Months of management", 44, "=C{months}", "months", "Programme plus procurement and closeout at both ends.", fmt="0", kind="input")
line("mgcost", "Client management cost", "=C{mgfte}*C{mgrate}*C{months}", "=D{mgfte}*D{mgrate}*D{months}", "£",
     "The line that is always paid and rarely counted, because it sits in overhead rather than in the package.")

sec("WHAT THE INTERFACES COST")
line("ifval", "Average cost when an interface fails", 45000, "=C{ifval}", "£", "INDICATIVE. Craneage standing, a re-order, a fortnight of programme. Set it from your own experience — this is the number you actually know better than the model does.", kind="input")
line("ifprob", "Probability an unowned interface bites", 0.35, 0.10, "%", "Five packages: roughly a third of unowned interfaces cost money. Under a prime the interface still exists but somebody owns it, so the residual is what leaks past them.", fmt="0%", kind="input")
line("ifcost", "Expected interface cost", "=C{ifaces}*C{ifval}*C{ifprob}", "=C{ifaces}*C{ifval}*D{ifprob}", "£",
     "Note the prime column still uses the same interface count. Consolidation does not remove interfaces; it gives them an owner.")

sec("WHAT THE PRIME CHARGES FOR TAKING IT ON")
line("margin", "Prime margin on the packages", 0.00, 0.085, "%", "The premium for one throat to choke. Nobody carries interface risk for nothing. 8.5% is INDICATIVE — this is the number to negotiate and the number to test below.", fmt="0.0%", kind="input")
line("marcost", "Prime margin cost", "=(C{direct}+C{directop})*C{margin}", "=(D{direct}+D{directop})*D{margin}", "£", "")
line("proccost", "Procurement and tendering cost", 95000, 42000, "£", "Five ITTs, five evaluations, five negotiations, five contracts — against one. INDICATIVE.", kind="input")

sec("THE ANSWER")
line("total", "TOTAL COST TO THE CLIENT", "=C{direct}+C{directop}+C{mgcost}+C{ifcost}+C{marcost}+C{proccost}",
     "=D{direct}+D{directop}+D{mgcost}+D{ifcost}+D{marcost}+D{proccost}", "£", "", bold=True)
line("delta", "Difference", "", "=C{total}-D{total}", "£", "Positive means one prime is cheaper on these assumptions.", bold=True)
line("be", "Break-even prime margin", "", "=IFERROR((C{total}-D{direct}-D{directop}-D{mgcost}-D{ifcost}-D{proccost})/(D{direct}+D{directop}),0)", "%",
     "Above this margin, five packages win. Below it, the prime does. This is the single number to take into the negotiation — and the one worth knowing before the first conversation.", fmt="0.0%", bold=True)

r += 1
band(ws, r, "   WHAT THE MODEL DOES NOT PRICE", 6); r += 1
for t in [
 "Programme. A prime can sequence across packages; five contractors each protect their own float and none of them owns the critical path between them.",
 "Single point of accountability when it goes wrong at 02:00 on a Sunday. This has a value and it is not in any of the cells above.",
 "Competition. Five packages are competed five ways; a prime is competed once and then you are married to it for 38 months.",
 "Supply chain access. A prime with standing relationships can buy better than a client buying five packages cold — sometimes enough to fund its own margin.",
 "The consolidation route itself. Appointing five and then novating them under one is a third option, and it is the one that keeps the competition and buys the ownership. It is also the hardest to contract, which is why it is worth being the person who has done it.",
]:
    c = ws.cell(row=r, column=2, value="·  " + t); c.font = BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 28
    r += 1
wb.save("/home/user/etablix/tools/village-plan.xlsx")
print("Packaging sheet written; rows:", K)
